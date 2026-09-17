"""专题导出：按收藏顺序生成 Markdown / docx 研究报告。"""
from io import BytesIO

from docx import Document as DocxDocument
from docx.shared import Pt


def _load(conn, topic_id):
    topic = conn.execute(
        "SELECT id, name, description FROM topics WHERE id=?", (topic_id,)
    ).fetchone()
    if topic is None:
        return None
    items = conn.execute(
        "SELECT i.article_id, a.article_label AS label, a.branch, a.chapter, a.section,"
        " a.content, d.title AS doc_title, d.category"
        " FROM topic_items i JOIN articles a ON a.id=i.article_id"
        " JOIN documents d ON d.id=a.document_id"
        " WHERE i.topic_id=? ORDER BY i.order_index",
        (topic_id,),
    ).fetchall()
    notes = conn.execute(
        "SELECT n.id, n.article_id, n.content_md, n.updated_at, a.article_label"
        " FROM notes n LEFT JOIN articles a ON a.id=n.article_id"
        " WHERE n.topic_id=? ORDER BY n.created_at, n.id",
        (topic_id,),
    ).fetchall()
    return topic, items, notes


def _split_notes(notes):
    """按法条分组；未关联法条的归入专题级笔记。"""
    by_article, topic_notes = {}, []
    for n in notes:
        if n["article_id"]:
            by_article.setdefault(n["article_id"], []).append(n)
        else:
            topic_notes.append(n)
    return by_article, topic_notes


def _location(item):
    parts = [p for p in (item["branch"], item["chapter"], item["section"]) if p]
    return " › ".join(parts)


def export_markdown(conn, topic_id):
    data = _load(conn, topic_id)
    if data is None:
        return None
    topic, items, notes = data
    by_article, topic_notes = _split_notes(notes)

    lines = [f"# {topic['name']}", ""]
    if topic["description"]:
        lines += [f"> {topic['description']}", ""]

    for i, item in enumerate(items, 1):
        loc = _location(item)
        head = f"## {i}. {item['label']}（{item['doc_title']}）"
        lines += [head]
        if loc:
            lines += [f"**层级**：{loc}", ""]
        else:
            lines += [""]
        lines += [item["content"], ""]
        for n in by_article.get(item["article_id"], []):
            quoted = "\n".join(
                "> " + l if l else ">" for l in n["content_md"].splitlines()
            )
            lines += [f"> **笔记**（{n['updated_at']}）", quoted, ""]

    if topic_notes:
        lines += ["## 专题笔记", ""]
        for n in topic_notes:
            quoted = "\n".join(
                "> " + l if l else ">" for l in n["content_md"].splitlines()
            )
            lines += [f"> {n['updated_at']}", quoted, ""]

    lines += [f"---", f"*导出自 LexBench · 法研台 · 共 {len(items)} 条法条 / {len(notes)} 则笔记*"]
    return "\n".join(lines)


def export_docx(conn, topic_id):
    data = _load(conn, topic_id)
    if data is None:
        return None
    topic, items, notes = data
    by_article, topic_notes = _split_notes(notes)

    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    doc.add_heading(topic["name"], 0)
    if topic["description"]:
        p = doc.add_paragraph(topic["description"])
        p.runs[0].italic = True

    for i, item in enumerate(items, 1):
        doc.add_heading(f"{i}. {item['label']}（{item['doc_title']}）", level=2)
        loc = _location(item)
        if loc:
            p = doc.add_paragraph(f"层级：{loc}")
            p.runs[0].font.size = Pt(9)
        for para in item["content"].splitlines():
            if para.strip():
                doc.add_paragraph(para)
        for n in by_article.get(item["article_id"], []):
            for line in n["content_md"].splitlines():
                p = doc.add_paragraph(f"【笔记】{line}" if line else "")
                p.runs[0].italic = True

    if topic_notes:
        doc.add_heading("专题笔记", level=2)
        for n in topic_notes:
            for line in n["content_md"].splitlines():
                p = doc.add_paragraph(line)
                p.runs[0].italic = True

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
